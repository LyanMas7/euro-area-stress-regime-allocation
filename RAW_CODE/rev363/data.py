"""REV363 locked-data readers and PIT as-of monthly panel.

Only the single locked ASSET03 archive supplied with REV362 is read.  The
readers are deliberately explicit: every source value carries an
``available_at`` rule, and the panel is assembled with a strict
``available_at <= issue_month_end`` filter.  Quarterly carry is a carry of one
native observation, never a new observation.

Status of the release clocks: ``RESEARCH_ASOF_NOT_OFFICIAL_ISSUER_PIT``.  The
supplied CLIFS file has an authentic AvailabilityDate column; the other rules
are the REV362 W01 research policies, reproduced here so the REV363 path does
not silently change the clock.
"""
from __future__ import annotations

import calendar
import math
from datetime import date, timedelta
from pathlib import Path

import numpy as np
import pandas as pd

from .config import (EA7, PANEL_START, RELEASE_RULES, SOURCES7, WARMUP_START)


# ---------------------------------------------------------------------------
# path discovery / source identity
# ---------------------------------------------------------------------------
def find_dataset_root(start) -> Path:
    p = Path(start).resolve()
    for c in (p, *p.parents):
        if (c / "06_BENCHMARK" / "ESRB_RQ01_CANONICAL_BENCHMARK.xlsx").is_file():
            return c
    hits = [d for d in p.rglob("ESRB_RQ01_CANONICAL_BENCHMARK.xlsx")]
    if len(hits) == 1:
        return hits[0].parent.parent
    raise FileNotFoundError("EXACTLY_ONE_LOCKED_DATASET_ROOT_REQUIRED")


def _exact(root: Path, name: str) -> Path:
    hits = list(root.rglob(name))
    if len(hits) != 1:
        raise FileNotFoundError(f"EXACTLY_ONE_SOURCE_REQUIRED:{name}:{len(hits)}")
    return hits[0]


def _prefer(root: Path, names: tuple[str, ...]) -> Path:
    """Prefer the first EA7 file present; fall back to the original EA8 name."""
    for name in names:
        hits = list(root.rglob(name))
        if len(hits) == 1:
            return hits[0]
        if len(hits) > 1:
            raise FileNotFoundError(f"EXACTLY_ONE_SOURCE_REQUIRED:{name}:{len(hits)}")
    raise FileNotFoundError("REQUIRED_SOURCE_MISSING:" + ",".join(names))


def month_end(y: int, m: int) -> date:
    return date(y, m, calendar.monthrange(y, m)[1])


def month_range(start: str, end: str) -> list[str]:
    y, m = int(start[:4]), int(start[5:7])
    ey, em = int(end[:4]), int(end[5:7])
    out = []
    while (y, m) <= (ey, em):
        out.append(f"{y:04d}-{m:02d}")
        m += 1
        if m == 13:
            y, m = y + 1, 1
    return out


def add_months(month: str, k: int) -> str:
    y = int(month[:4])
    m = int(month[5:7]) - 1 + k
    y += m // 12
    m %= 12
    return f"{y:04d}-{m + 1:02d}"


def _asof(dates: list[date], values: list, cutoff: date):
    """Latest observation with date <= cutoff (dates ascending)."""
    lo, hi = 0, len(dates)
    while lo < hi:
        mid = (lo + hi) // 2
        if dates[mid] <= cutoff:
            lo = mid + 1
        else:
            hi = mid
    if lo == 0:
        return None, None
    i = lo - 1
    return dates[i], values[i]


def truncate_rows(dates: list, columns: list, end: date):
    """Physical air-gap: drop every row dated after `end`, keeping the column lists aligned.

    As-of lookups could never reach such rows, so this changes no DEV number; it only guarantees that nothing
    dated after `data_end` stays in memory once a source file has been read."""
    keep = [i for i, d in enumerate(dates) if d <= end]
    if len(keep) == len(dates):
        return list(dates), [list(c) for c in columns]
    return [dates[i] for i in keep], [[c[i] for i in keep] for c in columns]


# ---------------------------------------------------------------------------
# raw source readers
# ---------------------------------------------------------------------------
def read_ciss(root: Path) -> tuple[list[date], list[float]]:
    df = pd.read_csv(_exact(root, "ciss_total_weekly.csv"))
    df.columns = ["date", "period", "value"]
    df["date"] = pd.to_datetime(df["date"]).dt.date
    df = df.dropna(subset=["value"]).sort_values("date")
    return list(df["date"]), list(df["value"].astype(float))


def read_vix(root: Path):
    df = pd.read_csv(_exact(root, "VIX_Daily.csv"))
    df["observation_date"] = pd.to_datetime(df["observation_date"]).dt.date
    df = df.dropna(subset=["VIXCLS"]).sort_values("observation_date")
    return list(df["observation_date"]), list(df["VIXCLS"].astype(float))


def read_clifs(root: Path, countries=EA7):
    df = pd.read_csv(_prefer(root, ("CLIFS_EA7_Monthly.csv", "CLIFS_EA8_Monthly.csv")))
    df["Date"] = pd.to_datetime(df["Date"]).dt.date
    df["AvailabilityDate"] = pd.to_datetime(df["AvailabilityDate"]).dt.date
    df = df[df["Country"].isin(countries)].sort_values(["Country", "Date"])
    out = {}
    for c in countries:
        sub = df[df["Country"] == c]
        out[c] = (list(sub["Date"]), list(sub["CLIFS"].astype(float)),
                  list(sub["AvailabilityDate"]))
    return out


def read_gfc(root: Path):
    from openpyxl import load_workbook
    p = _exact(root, "GFC_Price_Construction.xlsx")
    wb = load_workbook(p, read_only=True, data_only=True)
    rows = []
    for row in wb["Raw_Data"].iter_rows(min_row=2, values_only=True):
        if row[2] is None:
            continue
        d = pd.to_datetime(row[2]).date()
        v = row[3]
        if v is None or not np.isfinite(v):
            v = None
        rows.append((d, None if v is None else float(v)))
    wb.close()
    rows.sort()
    # prior-only expanding z with 24-month warm-up; G = -z
    dates, vals, prior = [], [], []
    for d, v in rows:
        n = len(prior)
        z = None
        if v is not None and n >= 24:
            mean = float(np.mean(prior))
            sd = float(np.std(prior, ddof=1))
            if sd > 0:
                z = (v - mean) / sd
        dates.append(d)
        vals.append(None if z is None else -z)
        if v is not None:
            prior.append(v)
    return dates, vals


def read_cg(root: Path):
    df = pd.read_csv(_exact(root, "ECB_ESRB_Euro_Area_Credit_to_GDP_Gap.csv"))
    df.columns = ["date", "period", "value"]
    df["date"] = pd.to_datetime(df["date"]).dt.date
    df = df.dropna(subset=["value"]).sort_values("date")
    return list(df["date"]), list(df["value"].astype(float))


def read_dsr(root: Path, countries=EA7):
    p = _prefer(root, ("EA7_Debt_Service_Ratios.csv", "EA8_Debt_Service_Ratios.csv"))
    df = pd.read_csv(p, skiprows=4)
    df.columns = ["key", "freq", "country", "borrower", "unit", "mult", "period",
                  "conf", "prebreak", "status", "value"]
    df = df[df["borrower"].astype(str).str.startswith("P:")].copy()
    df["country"] = df["country"].astype(str).str[:2]
    df["period"] = pd.to_datetime(df["period"]).dt.date
    df["value"] = pd.to_numeric(df["value"], errors="coerce")
    df = df.dropna(subset=["value"]).sort_values(["country", "period"])
    return {c: (list(df[df["country"] == c]["period"]),
                list(df[df["country"] == c]["value"].astype(float))) for c in countries}


def read_dfc_inputs(root: Path):
    p = _prefer(root, ("DFC_Construction_EA7.xlsx", "DFC_Construction.xlsx"))
    x = pd.ExcelFile(p)
    tc = pd.read_excel(x, sheet_name="RAW_TC_SELECTED")
    rpp = pd.read_excel(x, sheet_name="RAW_RPP_SELECTED")
    cpi = pd.read_excel(x, sheet_name="RAW_CPI_MONTHLY")
    for d in (tc, rpp, cpi):
        d["Date"] = pd.to_datetime(d["Date"]).dt.date
    return tc, rpp, cpi


def cf_weights(low: int = 32, high: int = 120, window: int = 120) -> np.ndarray:
    a, b = 2 * math.pi / high, 2 * math.pi / low
    b0 = (b - a) / math.pi
    bj = [(math.sin(b * j) - math.sin(a * j)) / (math.pi * j) for j in range(1, window - 1)]
    by_lag = [0.5 * b0, *bj, -0.5 * b0 - math.fsum(bj)]
    return np.asarray(by_lag[::-1], dtype=float)


_CF_W = cf_weights()


def yoy_log(levels):
    out = [None] * len(levels)
    for i in range(4, len(levels)):
        a, b = levels[i], levels[i - 4]
        if a is not None and b is not None and a > 0 and b > 0:
            out[i] = 100.0 * (math.log(a) - math.log(b))
    return out


def prior_z(cycle, anchor: int = 123, min_prior: int = 8):
    out = [None] * len(cycle)
    for i in range(anchor, len(cycle)):
        if cycle[i] is None:
            continue
        pr = [c for c in cycle[anchor:i] if c is not None]
        if len(pr) < min_prior:
            continue
        m = float(np.mean(pr))
        sd = float(np.std(pr, ddof=1))
        if sd > 0:
            out[i] = (cycle[i] - m) / sd
    return out


def construct_dfc_country(tc, rpp, cpi, country: str, last_quarter: str = "2024-Q4"):
    """Signed DFC exactly as the REV362 W01 DFC contract, extended to 2024-Q4."""
    tc_c = tc[tc["Country"] == country].sort_values("Date")
    rpp_c = rpp[rpp["Country"] == country].sort_values("Date")
    cpi_c = cpi[cpi["Country"] == country].sort_values("Date")
    qs = []
    y, q = 1971, 2
    ey, eq = int(last_quarter[:4]), int(last_quarter[-1])
    while (y, q) <= (ey, eq):
        qs.append(f"{y}-Q{q}")
        q += 1
        if q == 5:
            y, q = y + 1, 1
    tcm = {(d.year, (d.month - 1) // 3 + 1): (n, r)
           for d, n, r in zip(tc_c["Date"], tc_c["Nominal_Private_Credit_XDC"],
                              tc_c["Credit_to_GDP_pct"])}
    rpm = {(d.year, (d.month - 1) // 3 + 1): v for d, v in zip(rpp_c["Date"], rpp_c["Real_RPP_Index"])}
    cpim = {(d.year, d.month): v for d, v in zip(cpi_c["Date"], cpi_c["CPI_Index"])}
    nominal, ratio, house, cpi_qe = [], [], [], []
    for qq in qs:
        yy, qn = int(qq[:4]), int(qq[-1])
        n, r = tcm.get((yy, qn), (None, None))
        nominal.append(n)
        ratio.append(r)
        house.append(rpm.get((yy, qn)))
        cpi_qe.append(cpim.get((yy, qn * 3)))
    real = [n * 100.0 / c if n is not None and c not in (None, 0) else None
            for n, c in zip(nominal, cpi_qe)]
    series = {"real_credit": yoy_log(real), "credit_to_gdp": ratio, "real_rpp": yoy_log(house)}
    zs = {}
    for k, vals in series.items():
        cyc = [None] * len(vals)
        for i in range(119, len(vals)):
            win = vals[i - 119:i + 1]
            if any(v is None for v in win):
                continue
            arr = np.asarray(win, dtype=float)
            drift = (arr[-1] - arr[0]) / 119.0
            cyc[i] = float(arr @ _CF_W - drift * (_CF_W @ np.arange(120)))
        zs[k] = prior_z(cyc)
    out = {}
    for i, qq in enumerate(qs):
        z = [zs[k][i] for k in ("real_credit", "credit_to_gdp", "real_rpp")]
        out[qq] = float(np.mean(z)) if all(v is not None for v in z) else None
    return out


# ---------------------------------------------------------------------------
# monthly raw snapshot
# ---------------------------------------------------------------------------
def build_monthly_raw(root, *, row_start: str = PANEL_START, row_end: str = "2024-12",
                      data_end: str = "2024-12", countries=EA7, descriptive_ie_audit: bool = False) -> pd.DataFrame:
    """Assemble monthly research raw for the requested country scope (default EA7).

    EA7 ONLY: any country outside BE,DE,ES,FI,FR,IT,NL (IE, PT, ...) is refused unless the caller passes the
    explicit `descriptive_ie_audit=True`, which exists solely for the isolated W01 raw-descriptive audit and is
    never authority to fit an EA8 model or invent native issuer predictions.

    ``data_end`` bounds the month grid AND the rows kept: every source file is opened whole by its reader,
    but rows dated after ``data_end`` are dropped right after reading (``truncate_rows``), so the DEV call
    (data_end = 2019-12) keeps no 2020+ value in memory and none can reach an as-of lookup.
    ``row_start``/``row_end`` select the returned model rows;
    pre-row_start history is retained for one-sided feature construction.
    """
    if not set(countries) <= set(EA7) and not descriptive_ie_audit:
        raise ValueError("EA7_ONLY: non-EA7 countries need descriptive_ie_audit=True (audit-only, never a fit): "
                         + ",".join(sorted(set(countries) - set(EA7))))
    root = Path(root)
    months = month_range(PANEL_START, data_end)
    m_end = {p: month_end(int(p[:4]), int(p[5:7])) for p in months}

    ciss_d, ciss_v = read_ciss(root)
    vix_d, vix_v = read_vix(root)
    clifs = read_clifs(root, countries=countries)
    gfc_d, gfc_v = read_gfc(root)
    cg_d, cg_v = read_cg(root)
    dsr = read_dsr(root, countries=countries)
    tc, rpp, cpi = read_dfc_inputs(root)
    end_d = m_end[months[-1]]
    ciss_d, (ciss_v,) = truncate_rows(ciss_d, [ciss_v], end_d)
    vix_d, (vix_v,) = truncate_rows(vix_d, [vix_v], end_d)
    gfc_d, (gfc_v,) = truncate_rows(gfc_d, [gfc_v], end_d)
    cg_d, (cg_v,) = truncate_rows(cg_d, [cg_v], end_d)
    clifs = {c: (lambda t: (t[0], t[1][0], t[1][1]))(truncate_rows(ds, [vs, av], end_d))
             for c, (ds, vs, av) in clifs.items()}
    dsr = {c: (lambda t: (t[0], t[1][0]))(truncate_rows(dd, [vv], end_d)) for c, (dd, vv) in dsr.items()}
    tc, rpp, cpi = (f[f["Date"] <= end_d] for f in (tc, rpp, cpi))
    last_q = f"{int(data_end[:4])}-Q{(int(data_end[5:7]) - 1) // 3 + 1}"
    dfc = {c: construct_dfc_country(tc, rpp, cpi, c, last_quarter=last_q) for c in countries}

    def snap(dates, values, p, lag_days):
        d, v = _asof(dates, values, m_end[p] - timedelta(days=lag_days))
        return d, v

    ciss_snap = [snap(ciss_d, ciss_v, p, RELEASE_RULES["CISS"]["lag_days"]) for p in months]
    vix_snap = [snap(vix_d, vix_v, p, RELEASE_RULES["VIX"]["lag_days"]) for p in months]
    gfc_snap = [snap(gfc_d, gfc_v, p, RELEASE_RULES["GFC"]["lag_days"]) for p in months]
    ciss_series = [v for _, v in ciss_snap]
    vix_series = [v for _, v in vix_snap]
    gfc_series = [v for _, v in gfc_snap]
    cg_series = [snap(cg_d, cg_v, p, RELEASE_RULES["CG"]["lag_days"])[1] for p in months]
    clifs_series, clifs_observed_at, clifs_available_at, dsr_series = {}, {}, {}, {}
    for c in countries:
        ds, vs, av = clifs[c]
        # availability is explicit; walk forward to the latest available row
        vals, obs_at, av_at = [], [], []
        for p in months:
            eligible = [(d, v, a) for d, v, a in zip(ds, vs, av) if a <= m_end[p]]
            if eligible:
                d, v, a = max(eligible, key=lambda x: (x[0], x[2]))
            else:
                d, v, a = None, None, None
            vals.append(v); obs_at.append(d); av_at.append(a)
        clifs_series[c] = vals
        clifs_observed_at[c], clifs_available_at[c] = obs_at, av_at
        dd, vv = dsr[c]
        dsr_series[c] = [snap(dd, vv, p, RELEASE_RULES["DSR-P"]["lag_days"])[1] for p in months]

    rows = []
    for c in countries:
        qdate = {qq: month_end(int(qq[:4]), int(qq[-1]) * 3) for qq in dfc[c]}
        sorted_q = sorted(qdate)
        for i, p in enumerate(months):
            # latest DFC quarter available by month-end
            avail = [qq for qq in sorted_q
                     if qdate[qq] + timedelta(days=RELEASE_RULES["DFC"]["lag_days"]) <= m_end[p]]
            dfc_period = avail[-1] if avail else None
            dfc_val = dfc[c][dfc_period] if dfc_period else None
            dfc_obs = qdate[dfc_period] if dfc_period else None
            rows.append({
                "country": c, "month": p,
                "ciss": ciss_series[i], "vix": vix_series[i], "gfc": gfc_series[i],
                "clifs": clifs_series[c][i], "dfc": dfc_val,
                "cg": cg_series[i], "dsr": dsr_series[c][i],
                "observed_at_CISS": ciss_snap[i][0],
                "observed_at_VIX": vix_snap[i][0],
                "observed_at_GFC": gfc_snap[i][0],
                "observed_at_CLIFS": clifs_observed_at[c][i],
                "observed_at_DFC": dfc_obs,
                "native_period_DFC": dfc_period,
                "available_at_CISS": (ciss_snap[i][0] + timedelta(days=RELEASE_RULES["CISS"]["lag_days"])) if ciss_snap[i][0] else None,
                "available_at_VIX": vix_snap[i][0],
                "available_at_GFC": (gfc_snap[i][0] + timedelta(days=RELEASE_RULES["GFC"]["lag_days"])) if gfc_snap[i][0] else None,
                "available_at_CLIFS": clifs_available_at[c][i],
                "available_at_DFC": (dfc_obs + timedelta(days=RELEASE_RULES["DFC"]["lag_days"])) if dfc_obs else None,
            })
    df = pd.DataFrame(rows)
    df = df[(df["month"] >= row_start) & (df["month"] <= row_end)].reset_index(drop=True)
    return df
